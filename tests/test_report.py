"""Test laporan Excel, provenance, MOT interface, dan evaluasi detector."""
from __future__ import annotations

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from conftest import AHP_MATRIX, CRITERIA, ORIENTATION
from src.ahp import ahp_weights
from src.evaluate_detector import (DetectorEvalUnavailable, evaluate_detector,
                                   evaluation_status, not_available_frame)
from src.mot_eval import GroundTruthUnavailable, evaluate_mot, mot_status
from src.provenance import NA, build_provenance, environment_info
from src.report import SHEET_ORDER, not_run, write_report
from src.saw import saw_score
from src.sensitivity import one_at_a_time


@pytest.fixture
def bundle():
    rng = np.random.default_rng(5)
    idx = [f"S{i:02d}" for i in range(1, 12)]
    m = pd.DataFrame(rng.integers(1, 90, size=(11, 4)), columns=CRITERIA,
                     index=idx).astype(float)
    m.index.name = "entity"
    ahp = ahp_weights(np.array(AHP_MATRIX), CRITERIA)
    saw = saw_score(m, ahp["weights"], ORIENTATION)
    sens = one_at_a_time(m, ahp["weights"], ORIENTATION)
    return m, ahp, saw, sens


def test_report_backward_compatible_signature(tmp_path, bundle):
    """Pemanggilan gaya baseline (posisional) harus tetap berfungsi."""
    m, ahp, saw, sens = bundle
    p = tmp_path / "r.xlsx"
    write_report(str(p), m, ahp, saw, sens, {"a": 1})
    assert p.exists()


def test_report_contains_all_target_sheets(tmp_path, bundle):
    m, ahp, saw, sens = bundle
    p = tmp_path / "r.xlsx"
    write_report(str(p), m, ahp, saw, sens, {"cfg": {"nested": 1}})
    names = load_workbook(p).sheetnames
    for s in SHEET_ORDER:
        assert s in names, f"sheet hilang: {s}"
    assert len(names) >= 19


def test_missing_experiments_marked_not_run(tmp_path, bundle):
    m, ahp, saw, sens = bundle
    p = tmp_path / "r.xlsx"
    write_report(str(p), m, ahp, saw, sens, {"a": 1})
    wb = load_workbook(p)
    for sheet in ["Sensitivity_MonteCarlo", "Ablation", "AHP_SAW_vs_WP",
                  "Detection_Metrics", "Temporal_Comparison"]:
        vals = [c.value for row in wb[sheet].iter_rows(min_row=1, max_row=2)
                for c in row]
        assert "NOT_RUN" in vals, f"{sheet} harus ditandai NOT_RUN"


def test_no_dummy_numbers_in_not_run_sheets(tmp_path, bundle):
    """Sheet yang tidak dijalankan tidak boleh memuat angka apa pun."""
    m, ahp, saw, sens = bundle
    p = tmp_path / "r.xlsx"
    write_report(str(p), m, ahp, saw, sens, {"a": 1})
    wb = load_workbook(p)
    for row in wb["Ablation"].iter_rows(min_row=2):
        for c in row:
            assert not isinstance(c.value, (int, float)), \
                "sheet NOT_RUN tidak boleh memuat angka dummy"


def test_test_only_marking(tmp_path, bundle):
    m, ahp, saw, sens = bundle
    p = tmp_path / "r.xlsx"
    write_report(str(p), m, ahp, saw, sens, {"a": 1}, test_only=True)
    wb = load_workbook(p)
    assert "TEST_ONLY_NOTICE" in wb.sheetnames
    assert any(n.endswith("_TESTONLY") for n in wb.sheetnames)
    assert "TEST ONLY" in str(wb["TEST_ONLY_NOTICE"]["A2"].value)


def test_report_formatting_applied(tmp_path, bundle):
    m, ahp, saw, sens = bundle
    p = tmp_path / "r.xlsx"
    write_report(str(p), m, ahp, saw, sens, {"a": 1})
    ws = load_workbook(p)["SAW_Normalized"]
    assert ws.freeze_panes is not None
    assert ws.auto_filter.ref is not None
    assert ws["A1"].font.bold


def test_ahp_provenance_in_consistency_sheet(tmp_path, bundle):
    m, ahp, saw, sens = bundle
    p = tmp_path / "r.xlsx"
    write_report(str(p), m, ahp, saw, sens, {"a": 1})
    ws = load_workbook(p)["AHP_Consistency"]
    metrics = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert "ahp_source" in metrics and "expert_count" in metrics


def test_empty_report_still_writes(tmp_path):
    p = tmp_path / "r.xlsx"
    write_report(str(p), config_snapshot={"a": 1})
    wb = load_workbook(p)
    assert "SAW_Ranking" in wb.sheetnames


def test_not_run_frame_shape():
    d = not_run("alasan", "Eksperimen")
    assert d.iloc[0]["status"] == "NOT_RUN"
    assert d.iloc[0]["value"] == "NOT_AVAILABLE"


# ------------------------------------------------------------- provenance
def test_provenance_reports_not_available_not_fabricated():
    cfg = {"source": {"model_path": "weights/best.pt", "video_path": "missing.mp4"},
           "detection": {"conf": 0.25}, "ahp": {"source": "predefined_model"},
           "saw": {"zero_epsilon": 1e-6}, "monte_carlo": {"seed": 42}}
    df = build_provenance(cfg)
    d = dict(zip(df["key"], df["value"]))
    assert d["video_fps"] == NA          # video tidak ada -> tidak ditebak
    assert d["python_version"]
    assert d["monte_carlo_seed"] == "42"


def test_environment_info_keys():
    e = environment_info()
    for k in ["python_version", "os", "torch_version", "ultralytics_version",
              "cuda_version", "gpu", "git_hash", "timestamp_utc"]:
        assert k in e


# ------------------------------------------------------------- MOT
def test_mot_never_fabricates_metrics():
    s = mot_status(None)
    assert s["HOTA"] == "NOT_AVAILABLE"
    assert s["IDF1"] == "NOT_AVAILABLE"
    assert s["MOTA"] == "NOT_AVAILABLE"
    assert s["status"] == "NOT_AVAILABLE"
    assert not s["ground_truth_available"]


def test_mot_raises_without_ground_truth():
    with pytest.raises(GroundTruthUnavailable, match="ground truth"):
        evaluate_mot(None, None)


def test_mot_validates_ground_truth_columns(tmp_path):
    p = tmp_path / "gt.csv"
    pd.DataFrame({"frame": [1], "wrong": [1]}).to_csv(p, index=False)
    with pytest.raises(ValueError, match="kehilangan kolom"):
        evaluate_mot(str(p), None)


def test_mot_ready_when_ground_truth_present(tmp_path):
    p = tmp_path / "gt.csv"
    pd.DataFrame({"frame": [1, 2], "gt_id": [1, 1], "x1": [0, 0], "y1": [0, 0],
                  "x2": [5, 5], "y2": [5, 5]}).to_csv(p, index=False)
    r = evaluate_mot(str(p), None)
    assert r["status"] == "READY_TO_COMPUTE"
    assert r["n_gt_ids"] == 1
    assert "TRACKEVAL" in r["HOTA"]     # tetap tidak mengarang angka


# ------------------------------------------------------------- detector eval
def test_detector_eval_unavailable_without_inputs():
    st = evaluation_status(None, None)
    assert not st["available"] and len(st["reasons"]) == 2
    with pytest.raises(DetectorEvalUnavailable, match="tidak dapat dijalankan"):
        evaluate_detector("missing.pt", "missing.yaml")


def test_detector_not_available_frame_has_no_numbers():
    d = not_available_frame(["tidak ada data.yaml"])
    assert d.iloc[0]["precision"] == "NOT_AVAILABLE"
    assert d.iloc[0]["mAP50"] == "NOT_AVAILABLE"
    assert d.iloc[0]["status"] == "NOT_RUN"
