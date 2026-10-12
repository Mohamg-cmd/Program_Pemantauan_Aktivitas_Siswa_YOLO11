"""
Integration test — pipeline lengkap pada data SINTETIK.

PENTING: seluruh angka yang dihasilkan di sini adalah TEST ONLY untuk
verifikasi perangkat lunak. Bukan hasil penelitian dan tidak boleh dikutip.
Tidak membutuhkan GPU, YOLO, atau video.
"""
from __future__ import annotations

import numpy as np
import pandas as pd
import pytest

from conftest import AHP_MATRIX, CRITERIA, ORIENTATION, make_seat_grid
from src.aggregate import build_activity_matrix, build_both_representations
from src.ahp import ahp_weights
from src.identity import (IdentityConfig, assign_track_to_seat,
                          build_student_level_detections, load_seat_map,
                          merge_track_fragments, validate_seat_assignment)
from src.report import write_report
from src.saw import saw_score, zero_diagnostics
from src.sensitivity import one_at_a_time

FPS = 30.0
N_STUDENTS = 21


@pytest.fixture
def synthetic(tmp_path):
    """21 siswa, masing-masing pecah menjadi 2-4 fragmen track."""
    rng = np.random.default_rng(7)
    seats_df = make_seat_grid(N_STUDENTS)
    p = tmp_path / "seats.csv"
    seats_df.to_csv(p, index=False)
    seats = load_seat_map(str(p), N_STUDENTS)

    rows, tid = [], 1
    for s in seats.itertuples():
        nfrag = int(rng.integers(2, 5))
        bounds = np.linspace(0, 600, nfrag + 1).astype(int)
        for k in range(nfrag):
            for f in range(bounds[k], bounds[k + 1]):
                cls = CRITERIA[int(rng.choice(4, p=[0.10, 0.15, 0.65, 0.10]))]
                dy = -90 if cls == "Berdiri" else 0     # berdiri: bbox memanjang
                cx, by = s.seat_x + rng.normal(0, 5), s.seat_y + rng.normal(0, 5)
                rows.append({"frame": f, "track_id": tid,
                             "class_id": CRITERIA.index(cls), "class_name": cls,
                             "confidence": float(rng.uniform(.5, .95)),
                             "x1": cx - 40, "y1": by - 120 + dy,
                             "x2": cx + 40, "y2": by})
            tid += 1
    return seats, pd.DataFrame(rows)


def test_fragments_collapse_to_expected_student_count(synthetic):
    """Banyak track mentah harus menyatu menjadi TEPAT 21 identitas siswa."""
    seats, det = synthetic
    cfg = IdentityConfig(expected_students=N_STUDENTS)
    a = assign_track_to_seat(det, seats, cfg)
    rep = validate_seat_assignment(a, seats, cfg)

    assert rep["passed"]
    assert rep["n_unique_track_ids"] > N_STUDENTS, "fixture harus terfragmentasi"
    assert rep["n_assigned_students"] == N_STUDENTS
    assert rep["n_track_fragments_merged"] == rep["n_unique_track_ids"] - N_STUDENTS


def test_student_matrix_never_exceeds_expected(synthetic):
    seats, det = synthetic
    cfg = IdentityConfig(expected_students=N_STUDENTS)
    sl = build_student_level_detections(assign_track_to_seat(det, seats, cfg))
    m = build_activity_matrix("", FPS, "", detections_df=sl,
                              entity_column="student_id", classes=CRITERIA,
                              min_track_frames=30)
    assert len(m) <= N_STUDENTS
    assert set(m.index) <= {f"S{i:02d}" for i in range(1, N_STUDENTS + 1)}
    assert not any(str(i).startswith("track_") for i in m.index)


def test_track_level_entities_never_look_like_students(synthetic):
    """Regresi: mode track_level tidak boleh menghasilkan entitas bergaya siswa."""
    _, det = synthetic
    m = build_activity_matrix("", FPS, "", detections_df=det,
                              entity_column="track_id", classes=CRITERIA,
                              min_track_frames=30)
    assert all(str(i).startswith("track_") for i in m.index)
    assert len(m) > N_STUDENTS      # track mentah lebih banyak daripada siswa


def test_full_pipeline_to_excel(synthetic, tmp_path):
    """Jalur lengkap: identitas -> agregasi -> AHP -> SAW -> sensitivity -> Excel."""
    seats, det = synthetic
    cfg = IdentityConfig(expected_students=N_STUDENTS)
    a = assign_track_to_seat(det, seats, cfg)
    assert validate_seat_assignment(a, seats, cfg)["passed"]

    t2s, audit = merge_track_fragments(a, FPS, cfg)
    sl = build_student_level_detections(a)
    both = build_both_representations(sl, FPS, classes=CRITERIA, min_track_frames=30)
    m = both["duration"]

    ahp = ahp_weights(np.array(AHP_MATRIX), CRITERIA)
    assert ahp["consistent"]
    saw = saw_score(m, ahp["weights"], ORIENTATION)
    sens = one_at_a_time(m, ahp["weights"], ORIENTATION)

    assert len(saw) == N_STUDENTS
    assert saw["rank"].min() == 1 and saw["rank"].max() <= N_STUDENTS
    assert np.isfinite(saw["V"]).all()
    assert len(sens) == 4 * 6

    out = tmp_path / "report.xlsx"
    write_report(str(out), activity_duration=both["duration"],
                 activity_frequency=both["frequency"], ahp_result=ahp,
                 saw_result=saw, sensitivity_df=sens, config_snapshot={"x": 1},
                 track_to_student=t2s, student_identity_audit=audit,
                 identity_assignments=a.head(500),
                 zero_handling=zero_diagnostics(m, ORIENTATION),
                 test_only=True)
    assert out.exists()

    from openpyxl import load_workbook
    assert "TEST_ONLY_NOTICE" in load_workbook(out).sheetnames


def test_total_frames_conserved(synthetic):
    """Tidak ada frame yang hilang atau digandakan saat penggabungan fragmen."""
    seats, det = synthetic
    cfg = IdentityConfig(expected_students=N_STUDENTS)
    a = assign_track_to_seat(det, seats, cfg)
    _, audit = merge_track_fragments(a, FPS, cfg)
    assigned = a[a.assignment_status == "assigned"]
    expected = assigned.groupby("student_id")["frame"].nunique().sum()
    assert audit["total_frames"].sum() == expected


def test_duration_frequency_consistency(synthetic):
    seats, det = synthetic
    cfg = IdentityConfig(expected_students=N_STUDENTS)
    sl = build_student_level_detections(assign_track_to_seat(det, seats, cfg))
    both = build_both_representations(sl, FPS, classes=CRITERIA, min_track_frames=30)
    pd.testing.assert_frame_equal(
        both["duration"], both["frequency"] / FPS, check_exact=False)
